import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc = Document()

    # Configurar margenes
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Titulo principal
    title = doc.add_heading("Guía Práctica y Explicativa del Taller 6", level=0)
    title.runs[0].font.color.rgb = RGBColor(14, 116, 144)

    sub = doc.add_paragraph("Sprint Review, Retrospective y Liberación Segura en Continuous Deployment (CD + TBD)")
    sub.runs[0].font.size = Pt(13)
    sub.runs[0].font.italic = True
    sub.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph(
        "Nota informal para el equipo: Esta guía resume de manera directa, sencilla y sin rodeos todo lo que se "
        "implementó y acordó durante el Taller 6, incluyendo las capturas de pantalla de los momentos más relevantes."
    )

    doc.add_heading("1. ¿De qué va este Taller 6? (En cristiano)", level=1)
    doc.add_paragraph(
        "Hasta el Taller 5 teníamos el pipeline completo: el código se prueba con CI, se integra a main y se despliega solo a Render, "
        "usando Feature Flags con ConfigCat para no romperle nada al usuario. Pero surgió la gran pregunta: si desplegamos varias veces por "
        "semana... ¿cómo cambian el Sprint Review y la Retrospective? En este taller adaptamos las ceremonias a este nuevo ritmo y practicamos "
        "la liberación segura en vivo resolviendo el caso de estudio de 'Pagos Express'."
    )

    doc.add_heading("2. La Nueva Sprint Review (Adiós a las demos en localhost)", level=1)
    doc.add_paragraph(
        "En Continuous Deployment, la Review ya no es para mostrar diapositivas ni código en máquinas locales. La demo se hace en producción (Render). "
        "Se muestra lo que ya está vivo o protegido por Feature Toggles, se exhibe telemetría real y se valida el valor directamente con el Product Owner."
    )
    p_format = doc.add_paragraph()
    p_format.add_run("Estructura obligatoria de demo (4-5 min por ítem):\n").bold = True
    p_format.add_run("1. ¿Qué problema de usuario resuelve?\n2. ¿Está activo para todos o detrás de toggle?\n3. Demo en vivo en Render (con toggle ON/OFF).\n4. Métricas de telemetría.\n5. Pregunta al PO: '¿Esto genera el valor esperado?'")

    if os.path.exists("capturas/captura_03_sprint_review_demos.png"):
        doc.add_paragraph().add_run().add_picture("capturas/captura_03_sprint_review_demos.png", width=Inches(6.2))
        cap = doc.add_paragraph("Figura 1: Tablero interactivo de Sprint Review con comparación oficial y 3 demos reales.")
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.italic = True

    doc.add_heading("3. Liberación Segura de Valor y el Caso 'Pagos Express'", level=1)
    doc.add_paragraph(
        "El caso presentaba la funcionalidad 'Pagos Express con un clic' en main hace 9 días detrás del toggle 'pagos-express-v1' al 0%. "
        "Había miedo por caídas previas en health check y falta de métricas. Se acordó la estrategia de liberación controlada:"
    )
    doc.add_paragraph("• Cohorte Canary (20%): Activación para usuarios beta con telemetría activa antes de abrir al 100% comercial.")
    doc.add_paragraph("• Kill Switch (< 1 segundo): Si la tasa de error supera el 1% o la latencia pasa de 500ms, el toggle se apaga inmediatamente en ConfigCat sin re-desplegar.")
    doc.add_paragraph("• Limpieza de Deuda: Identificación y remoción de 3 toggles viejos (> 45 días).")

    if os.path.exists("capturas/captura_04_pagos_express_canary_20.png"):
        doc.add_paragraph().add_run().add_picture("capturas/captura_04_pagos_express_canary_20.png", width=Inches(6.2))
        cap = doc.add_paragraph("Figura 2: Consola de Pago Express con cohorte Canary (20%) y telemetría en tiempo real.")
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.italic = True

    if os.path.exists("capturas/captura_05_kill_switch_emergencia.png"):
        doc.add_paragraph().add_run().add_picture("capturas/captura_05_kill_switch_emergencia.png", width=Inches(6.2))
        cap = doc.add_paragraph("Figura 3: Activación del Kill Switch de emergencia bloqueando la funcionalidad en < 1s.")
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.italic = True

    doc.add_heading("4. La Nueva Sprint Retrospective (Enfocada en el Pipeline)", level=1)
    doc.add_paragraph(
        "La Retrospectiva ahora inspecciona formalmente tres pilares técnicos: Salud del Pipeline (fallos de CI, tiempo de deploy, rollbacks), "
        "Disciplina TBD (integraciones diarias a main, vida de ramas < 24h) y Flujo de Valor (Lead Time)."
    )
    doc.add_paragraph("Siguiendo la regla de oro, se definieron exactamente 3 acciones de mejora asignadas:")
    doc.add_paragraph("1. Acción #1 (Tech Lead): Retiro de flags a los 30 días posteriores a GA. Prohibido acumular toggles de más de 45 días.")
    doc.add_paragraph("2. Acción #2 (Dev Team): Observabilidad obligatoria (latencia y error) antes de autorizar cualquier rollout > 0%.")
    doc.add_paragraph("3. Acción #3 (PO & SM): Política de No-Releases en Viernes; liberaciones canary solo de lunes a jueves con 24-48h de observación.")

    if os.path.exists("capturas/captura_06_sprint_retro_tablero.png"):
        doc.add_paragraph().add_run().add_picture("capturas/captura_06_sprint_retro_tablero.png", width=Inches(6.2))
        cap = doc.add_paragraph("Figura 4: Tablero de Sprint Retrospective estilo Miro con las 4 columnas de análisis.")
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.italic = True

    doc.add_heading("5. Salud Técnica y Validación del Pipeline (CI/CD)", level=1)
    doc.add_paragraph(
        "Para garantizar que main se mantuviera siempre verde y desplegable, se ejecutó la suite completa de pruebas: "
        "79 tests en verde (100% pasando en 0.23 segundos) y linters limpios con Ruff."
    )

    if os.path.exists("capturas/captura_01_ci_cd_tests_ruff.png"):
        doc.add_paragraph().add_run().add_picture("capturas/captura_01_ci_cd_tests_ruff.png", width=Inches(6.2))
        cap = doc.add_paragraph("Figura 5: Terminal de ejecución de pruebas y verificación de Ruff linter.")
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.italic = True

    if os.path.exists("capturas/captura_02_health_check_render.png"):
        doc.add_paragraph().add_run().add_picture("capturas/captura_02_health_check_render.png", width=Inches(5.0))
        cap = doc.add_paragraph("Figura 6: Endpoint /health verificado y operativo para Render.")
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.italic = True

    doc.add_heading("6. Playbook Versión 3.0 y Checklist Fase 1 al 100%", level=1)
    doc.add_paragraph(
        "Se actualizó SCRUM_TBD_PLAYBOOK.md a la versión 3.0 con todos los acuerdos oficiales y se completó al 100% el Checklist de Fase 1."
    )

    if os.path.exists("capturas/captura_07_checklist_fase1_playbook.png"):
        doc.add_paragraph().add_run().add_picture("capturas/captura_07_checklist_fase1_playbook.png", width=Inches(6.2))
        cap = doc.add_paragraph("Figura 7: Checklist Oficial de Fase 1 cerrado y validado al 100%.")
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.italic = True

    doc.add_heading("7. Conclusión del Taller", level=1)
    doc.add_paragraph(
        "\"Nunca más llegaremos a una Review a mostrar diapositivas o código local sin desplegar: mostramos lo que ya corre en producción "
        "tras toggles con métricas reales. Y en cada Retro, antes de hablar de percepciones, auditaremos primero la salud del pipeline, el tiempo de "
        "los deploys y la disciplina de integrar a main cada día.\""
    )

    doc.save("GUIA_EXPLICATIVA_TALLER_6.docx")
    print("Documento GUIA_EXPLICATIVA_TALLER_6.docx generado exitosamente.")

if __name__ == "__main__":
    main()
